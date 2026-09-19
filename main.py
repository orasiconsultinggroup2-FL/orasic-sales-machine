import os
import logging
import requests
import json
import io
import urllib.parse
import unicodedata
import re
from datetime import datetime

from fastapi import FastAPI, Request, Form
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, RedirectResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import uvicorn

# Configuración de Logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI()

# --- CONFIGURACIÓN DE ENTORNO ---
SUPABASE_URL = os.environ.get("SUPABASE_URL", "").rstrip("/")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY", "")
GOOGLE_MAPS_API_KEY = os.environ.get("GOOGLE_MAPS_API_KEY", "")

supabase_connected = False
if SUPABASE_URL and SUPABASE_KEY:
    try:
        headers = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}", "Content-Type": "application/json"}
        test_url = f"{SUPABASE_URL}/rest/v1/leads_maestros?select=id&limit=1"
        response = requests.get(test_url, headers=headers, timeout=5)
        if response.status_code == 200:
            supabase_connected = True
            logger.info("Conexión a Supabase EXITOSA")
    except Exception as e:
        logger.error(f"Error conexión Supabase: {e}")

# --- CONSTANTES ---
TERMINOS_RELACIONADOS = {
    "Barbería": ["Barbería", "Barber Shop", "Barbershop", "Peluquería masculina"],
    "Clínica Dental": ["Clínica Dental", "Dentista", "Odontólogo", "Consultorio dental"],
    "Veterinaria": ["Veterinaria", "Clínica veterinaria", "Pet Shop", "Tienda de mascotas"],
    "Spa": ["Spa", "Masajes", "Masoterapia", "Relajación"],
    "Gimnasio": ["Gimnasio", "Gym", "Fitness", "CrossFit"],
    "Cancha de Fútbol": ["Cancha de Fútbol", "Fútbol 5", "Complejo deportivo"],
    "Salón de Belleza": ["Salón de Belleza", "Peluquería", "Estética", "Manicure"],
    "Pilates": ["Pilates", "Estudio de Pilates", "Yoga"],
    "Clínica Médica": ["Clínica Médica", "Centro Médico", "Policlínico"],
    "Agencia de Marketing": ["Agencia de Marketing", "Marketing Digital", "Publicidad"],
    "Academia de Inglés": ["Academia de Inglés", "Instituto de idiomas", "Británico", "ICPNA"],
    "Panadería": ["Panadería", "Pastelería", "Bakery", "Pan artesanal"],
    "Zapaterías": ["Zapaterías", "Calzado", "Tienda de zapatos"]
}

RUBROS_NEGOCIOS = sorted(list(TERMINOS_RELACIONADOS.keys())) + ["Otro"]
DISTRITOS_LIMA = [
    "Ancón", "Ate", "Barranco", "Breña", "Carabayllo", "Chaclacayo", "Chorrillos", 
    "Cieneguilla", "Comas", "El Agustino", "Independencia", "Jesús María", "La Molina", 
    "La Victoria", "Lima", "Lince", "Los Olivos", "Lurigancho", "Lurín", 
    "Magdalena del Mar", "Miraflores", "Pachacámac", "Pucusana", "Pueblo Libre", 
    "Puente Piedra", "Punta Hermosa", "Punta Negra", "Rímac", "San Bartolo", 
    "San Borja", "San Isidro", "San Juan de Lurigancho", "San Juan de Miraflores", 
    "San Luis", "San Martín de Porres", "San Miguel", "Santa Anita", 
    "Santa María del Mar", "Santa Rosa", "Santiago de Surco", "Surquillo", 
    "Villa El Salvador", "Villa María del Triunfo", "Otro"
]

# --- FUNCIONES AUXILIARES ---

def normalizar_agresivo(texto):
    if not texto: return ""
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', texto.lower().strip())

def limpiar_numero(valor):
    if not valor and valor != 0: return ""
    try:
        num = float(str(valor).replace(',', ''))
        if num == int(num): return str(int(num))
        return str(num)
    except: return str(valor)

def generar_pitch_inteligente(nombre, rubro, distrito, rating, reseñas, web, telefono, plan):
    if not nombre: nombre = "Cliente Potencial"
    nombre_corto = str(nombre).split("–")[0].split("-")[0].strip()
    rubro_lower = (rubro or "").lower()
    
    if any(x in rubro_lower for x in ["dental", "médica", "veterinaria"]):
        angulo = "captar pacientes de alto valor mediante posicionamiento digital estratégico"
    elif any(x in rubro_lower for x in ["barber", "salón", "spa", "belleza", "zapato", "calzado"]):
        angulo = "llenar su agenda/tienda con clientes recurrentes mediante un sistema digital profesional"
    elif any(x in rubro_lower for x in ["gimnasio", "cancha", "academia", "pilates", "inglés"]):
        angulo = "maximizar la ocupación de sus instalaciones con automatización de reservas y captación digital"
    elif "marketing" in rubro_lower:
        angulo = "potenciar su cartera de clientes con herramientas de prospección automatizada"
    else:
        angulo = "potenciar sus ventas con una estrategia digital personalizada"

    if plan == "CUSTOM":
        return (f"Hola equipo de {nombre_corto}, he revisado su presencia digital y veo que ya tienen una estructura sólida. "
                f"Dado su nivel de operación, no tenemos un paquete estándar, pero sí solemos trabajar de dos formas con empresas de su talla: "
                f"1) Realizando una auditoría profunda de su embudo de ventas actual para detectar fugas de dinero ocultas, o "
                f"2) Integrando nuestro motor de IA directamente con su CRM actual para predecir la demanda estacional. "
                f"¿Les haría sentido agendar una llamada de 15 minutos para explorar cuál de estas dos vías les traería más ROI este trimestre?")
    elif plan == "PRO":
        return (f"Hola equipo de {nombre_corto}, vi que tienen {limpiar_numero(rating)} estrellas con {limpiar_numero(reseñas)} reseñas y una presencia digital consolidada. "
                f"Con ese volumen de demanda, usualmente ayudamos a marcas como la suya en dos cosas clave: "
                f"1) Implementando IA para atender consultas básicas 24/7 y liberar a su equipo de recepción, o "
                f"2) Analizando datos predictivos para saber exactamente qué servicios promocionar antes de que baje la ocupación. "
                f"¿Les gustaría ver cómo funciona la automatización con IA o prefieren que empecemos por la analítica de ocupación? Les mando un caso de éxito similar.")
    elif plan == "MANAGER":
        return (f"Hola equipo de {nombre_corto}, vi que tienen {limpiar_numero(rating)} estrellas con {limpiar_numero(reseñas)} reseñas en Google Maps. "
                f"Imagino que con ese volumen, coordinar la agenda debe ser un reto. En ORASIC Lab trabajamos con negocios de su nivel en dos frentes: "
                f"1) Automatizando las reservas para llenar los huecos libres sin intervención manual, o "
                f"2) Creando un sistema de fidelización para que sus clientes recurrentes vuelvan más seguido y gasten un poco más por visita. "
                f"¿Sienten que hoy les duele más la gestión del tiempo o la retención de clientes? Quedo atento para mostrarles cómo lo resolvemos.")
    else:
        return (f"Hola equipo de {nombre_corto}, vi que tienen {limpiar_numero(rating)} estrellas en {distrito} y un trato muy cercano que sus clientes valoran. "
                f"Con esa base de confianza, normalmente ayudamos a negocios como el suyo de dos formas: "
                f"1) Creando una página simple para que los nuevos clientes del barrio los encuentren fácil en Google, o "
                f"2) Implementando un sistema de reservas automático para que no pierdan tiempo coordinando citas por WhatsApp. "
                f"¿Alguna de estas dos opciones les resuena más ahora mismo? Me avisan y les envío una propuesta rápida.")

def get_all_leads():
    """Obtiene TODOS los leads de Supabase (Maestros + Google guardados)"""
    if not supabase_connected: return []
    try:
        headers = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}"}
        # Aumentamos límite a 3000 para asegurar traer todo
        url = f"{SUPABASE_URL}/rest/v1/leads_maestros?limit=3000&order=created_at.desc"
        resp = requests.get(url, headers=headers, timeout=15)
        return resp.json() if resp.status_code == 200 else []
    except Exception as e:
        logger.error(f"Error obteniendo leads: {e}")
        return []

def get_sent_count():
    if not supabase_connected: return 0
    try:
        headers = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}", "Prefer": "count=exact"}
        resp = requests.get(f"{SUPABASE_URL}/rest/v1/leads_maestros?Estado=eq.Enviado&select=id", headers=headers, timeout=5)
        range_header = resp.headers.get('Content-Range', '')
        if '/' in range_header: return int(range_header.split('/')[-1])
        return 0
    except: return 0

def insert_lead_supabase(lead):
    if not supabase_connected: return False
    try:
        clean_lead = lead.copy()
        # Limpieza estricta de tipos para evitar errores de inserción
        if 'Reseñas' in clean_lead:
            try: clean_lead['Reseñas'] = int(float(str(clean_lead['Reseñas']).replace('.0','').replace(',','')))
            except: clean_lead['Reseñas'] = 0
        if 'Rating' in clean_lead:
            try: clean_lead['Rating'] = float(str(clean_lead['Rating']).replace(',','.'))
            except: clean_lead['Rating'] = 0.0
            
        headers = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}", "Content-Type": "application/json", "Prefer": "return=minimal"}
        resp = requests.post(f"{SUPABASE_URL}/rest/v1/leads_maestros", headers=headers, json=clean_lead, timeout=5)
        if resp.status_code >= 400: logger.error(f"Error insertando: {resp.text[:100]}")
        return 200 <= resp.status_code < 300
    except Exception as e:
        logger.error(f"Excepción insertando: {e}")
        return False

def calificar_lead_google(reviews_str, rating_str, website):
    plan = "STARTER"
    has_web = any(x in (website or "").lower() for x in ['.pe', '.com', '.net'])
    
    try: rev_int = int(float(str(reviews_str).replace('.0', '').replace(',', '')))
    except: rev_int = 0
    try: rat_float = float(str(rating_str).replace(',', '.'))
    except: rat_float = 0.0
    
    if has_web and ('booking' in (website or "").lower() or 'fresha' in (website or "").lower()) and rev_int > 500 and rat_float >= 4.8:
        return "CUSTOM"
    if has_web and 100 <= rev_int <= 500 and rat_float >= 4.5: return "PRO"
    if has_web and rev_int > 500 and rat_float >= 4.7: return "PRO"
    # Lógica Manager ajustada: Exige 50+ reseñas Y (Web O Rating alto)
    if rev_int >= 50 and (has_web or rat_float >= 4.5): return "MANAGER"
    return "STARTER"

def search_and_save_leads(keyword, location, limit=100):
    """Busca en Google, CALIFICA, GUARDA EN BD y retorna los leads"""
    if not GOOGLE_MAPS_API_KEY: return [], 0
    
    terminos_busqueda = [keyword]
    keyword_clean = keyword.strip()
    for categoria, sinonimos in TERMINOS_RELACIONADOS.items():
        if categoria.lower() == keyword_clean.lower():
            terminos_busqueda = sinonimos
            break
    
    terminos_unicos = list(dict.fromkeys(terminos_busqueda))
    all_leads = []
    existing_leads = get_all_leads() # Traer existentes para evitar duplicados
    
    existing_keys = set()
    for l in existing_leads:
        n = normalizar_agresivo(l.get('Negocio', l.get('nombre', '')))
        d = normalizar_agresivo(l.get('Distrito', l.get('distrito', '')))
        existing_keys.add((n, d))
        
    count_saved = 0
    for termino in terminos_unicos:
        url = "https://places.googleapis.com/v1/places:searchText"
        headers = {
            "Content-Type": "application/json", 
            "X-Goog-Api-Key": GOOGLE_MAPS_API_KEY, 
            "X-Goog-FieldMask": "places.id,places.displayName,places.formattedAddress,places.primaryType,places.internationalPhoneNumber,places.rating,places.userRatingCount,places.websiteUri"
        }
        payload = {"textQuery": f"{termino} en {location}", "maxResultCount": min(limit, 50), "languageCode": "es"}
        
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=12)
            if resp.status_code != 200: continue
            places = resp.json().get("places", [])
            
            for place in places:
                name = place.get("displayName", {}).get("text", "").strip()
                if not name: continue
                
                name_norm = normalizar_agresivo(name)
                loc_norm = normalizar_agresivo(location)
                current_key = (name_norm, loc_norm)
                
                if current_key in existing_keys: continue # Saltar duplicados
                
                tel = place.get("internationalPhoneNumber", "").replace(" ", "").replace("-", "")
                rating = place.get("rating", "")
                reviews = place.get("userRatingCount", "")
                website = place.get("websiteUri", "")
                address = place.get("formattedAddress", "")
                
                presencia_digital = "Fuerte" if website else ("Media" if tel else "Baja")
                plan = calificar_lead_google(reviews, rating, website)
                pitch = generar_pitch_inteligente(name, keyword_clean, location, str(rating), str(reviews), website, tel, plan)
                
                lead = {
                    "Categoria": keyword_clean, "Distrito": location.title(), "Negocio": name,
                    "Direccion": address, "Rating": rating, "Reseñas": reviews,
                    "Contacto": tel, "Web/Redes": website,
                    "Publico objetivo": "", "Puntos fuertes": "", "Puntos debiles": "",
                    "Presencia digital": presencia_digital, "Responde reseñas": "",
                    "Plan_sugerido": plan, "pitch_automatizado": pitch,
                    "origen": "GOOGLE_SEARCH", "created_at": datetime.now().isoformat()
                }
                
                # GUARDAR INMEDIATAMENTE
                if insert_lead_supabase(lead):
                    count_saved += 1
                    all_leads.append(lead)
                    existing_keys.add(current_key) # Agregar a set local para evitar duplicados en misma búsqueda
                    
                if len(all_leads) >= limit: break
            if len(all_leads) >= limit: break
        except Exception as e:
            logger.error(f"Error Google: {e}")
            continue
            
    return all_leads, count_saved

# --- ENDPOINTS ---

@app.post("/api/auto-search")
async def auto_search(request: Request):
    form = await request.form()
    keyword = form.get("keyword", "")
    location = form.get("location", "")
    limit = min(int(form.get("limit", 100)), 100)
    
    if keyword == "Otro": keyword = form.get("keyword_other", "").strip()
    if location == "Otro": location = form.get("location_other", "").strip()

    if not keyword or not location:
        return RedirectResponse(url="/?error=Faltan+datos", status_code=303)

    # Buscar, Calificar y GUARDAR en BD
    new_leads, count_saved = search_and_save_leads(keyword, location, limit)
    
    msg = f"Busqué '{keyword}' en '{location}'. {len(new_leads)} encontrados. {count_saved} guardados en Base de Datos."
    
    # Redirigir SIN pasar los resultados por URL (ahora están en BD)
    redirect_url = (
        f"/?success={urllib.parse.quote(msg)}"
        f"&last_keyword={urllib.parse.quote(keyword)}"
        f"&last_location={urllib.parse.quote(location)}"
        f"&filter_origin=google" # Forzar vista Google al terminar búsqueda
    )
    return RedirectResponse(url=redirect_url, status_code=303)

@app.get("/", response_class=HTMLResponse)
async def dashboard(success: str = None, error: str = None, filter_origin: str = None, last_keyword: str = None, last_location: str = None):
    leads_enviados = get_sent_count()
    all_leads = get_all_leads() # LEER SIEMPRE DE BD, NUNCA DE URL
    
    # Separar leads por origen REAL en BD
    google_leads = [l for l in all_leads if l.get("origen") == "GOOGLE_SEARCH"]
    manual_leads = [l for l in all_leads if l.get("origen") != "GOOGLE_SEARCH"]
    
    # Determinar vista activa
    if filter_origin == "google":
        filtered_leads = google_leads
        view_title = "RESULTADOS GOOGLE"
    elif filter_origin == "manual":
        filtered_leads = manual_leads
        view_title = "BASE MANUAL (MAESTROS)"
    elif filter_origin == "all":
        filtered_leads = manual_leads + google_leads
        view_title = "TODOS (Base + Google)"
    else:
        # Por defecto mostrar Google si hay, sino Manual
        if google_leads:
            filtered_leads = google_leads
            view_title = "RESULTADOS DE BÚSQUEDA"
        else:
            filtered_leads = manual_leads
            view_title = "BASE MANUAL (MAESTROS)"
    
    alert_html = ""
    if success:
        alert_html = "<div style='background:#064E3B;border-left:4px solid #10B981;color:#ECFDF5;padding:15px;margin-bottom:20px;'>[OK] " + urllib.parse.unquote(success) + "</div>"
    elif error:
        alert_html = "<div style='background:#7F1D1D;border-left:4px solid #EF4444;color:#FEF2F2;padding:15px;margin-bottom:20px;'>[ERROR] " + urllib.parse.unquote(error) + "</div>"
    
    rows_html = ""
    for i, lead in enumerate(filtered_leads, 1):
        nombre = str(lead.get('Negocio', lead.get('nombre', 'Sin Nombre')))
        rubro = str(lead.get('Categoria', lead.get('rubro', '')))
        distrito = str(lead.get('Distrito', lead.get('distrito', '')))
        contacto = str(lead.get('Contacto', lead.get('telefono', '')))
        
        rating = limpiar_numero(lead.get('Rating', lead.get('rating', '')))
        reseñas = limpiar_numero(lead.get('Reseñas', lead.get('reseñas', '')))
        
        web = str(lead.get('Web/Redes', lead.get('web_redes', '')))
        plan = str(lead.get('Plan_sugerido', lead.get('plan_sugerido', 'STARTER')))
        origen = str(lead.get('origen', 'MANUAL_BASE'))
        publico_obj = str(lead.get('Publico objetivo', ''))
        puntos_f = str(lead.get('Puntos fuertes', ''))
        
        lead_id = lead.get('id', f"temp_{i}")
        
        pitch = lead.get('pitch_automatizado', '')
        if not pitch:
            pitch = generar_pitch_inteligente(nombre, rubro, distrito, rating, reseñas, web, contacto, plan)
        
        texto_msg = urllib.parse.quote(pitch)
        
        if origen == "GOOGLE_SEARCH":
            origen_badge = '<span style="background:#3B82F620;color:#3B82F6;padding:2px 8px;border-radius:4px;font-size:0.7rem;">[GOOGLE]</span>'
        else:
            origen_badge = '<span style="background:#F59E0B20;color:#F59E0B;padding:2px 8px;border-radius:4px;font-size:0.7rem;">[MAESTRO]</span>'
        
        btn_accion = ""
        if contacto and contacto != "":
            onclick_js = f"fetch('/api/mark-sent/{lead_id}',{{method:'POST'}});" if str(lead_id).isdigit() or len(str(lead_id)) > 5 else ""
            btn_accion = (
                f"<a href='https://wa.me/{contacto}?text={texto_msg}' target='_blank' "
                f"onclick='{onclick_js}' "
                f"style='background:#22C55E;color:white;padding:6px 12px;border-radius:6px;text-decoration:none;display:inline-block;white-space:nowrap;'>"
                f"[WA]</a>"
            )
        else:
            btn_accion = (
                f"<button onclick=\"navigator.clipboard.writeText(decodeURIComponent('{texto_msg}'));this.textContent='[COPIADO]';setTimeout(()=>this.textContent='[COPIAR]',2000);\" "
                f"style='background:#64748B;color:white;padding:6px 12px;border:none;border-radius:6px;cursor:pointer;'>[COPIAR]</button>"
            )
        
        rating_display = f"{rating} ({reseñas})" if rating and reseñas and str(reseñas) != "0" else "-"
        web_display = f'<a href="{web}" target="_blank" style="color:#22D3EE;text-decoration:none;">[WEB]</a>' if web else "-"
        
        plan_color = "#10B981" if plan == "CUSTOM" else ("#8B5CF6" if plan == "PRO" else ("#3B82F6" if plan == "MANAGER" else "#64748B"))
            
        rows_html += (
            f"<tr style='border-bottom:1px solid #1E293B;'>"
            f"<td style='padding:12px;'>#{i}</td>"
            f"<td style='font-weight:600;'>{nombre}</td>"
            f"<td>{rubro}</td>"
            f"<td>{distrito}</td>"
            f"<td>{contacto if contacto else '-'}</td>"
            f"<td>{rating_display}</td>"
            f"<td>{web_display}</td>"
            f"<td><span style='background:{plan_color}20;color:{plan_color};padding:4px 10px;border-radius:12px;font-size:0.8rem;font-weight:bold;'>{plan}</span></td>"
            f"<td>{origen_badge}</td>"
            f"<td>{btn_accion}</td>"
            f"</tr>"
            f"<tr style='border-bottom:1px solid #1E293B; background:#0F172A;'>"
            f"<td colspan='10' style='padding:8px 12px; font-size:0.85rem; color:#94A3B8;'>"
            f"<strong>Público:</strong> {publico_obj or '-'} | "
            f"<strong>Fortalezas:</strong> {puntos_f or '-'}"
            f"</td>"
            f"</tr>"
        )
    
    if not rows_html: 
        rows_html = "<tr><td colspan='10' style='padding:30px;text-align:center;'>No hay prospectos en esta vista.</td></tr>"
    
    rubro_options = "".join([f'<option value="{r}" {"selected" if r == last_keyword else ""}>{r}</option>' for r in RUBROS_NEGOCIOS])
    distrito_options = "".join([f'<option value="{d}" {"selected" if d == last_location else ""}>{d}</option>' for d in DISTRITOS_LIMA])
    
    count_manual = len(manual_leads)
    count_google = len(google_leads)
    count_total = count_manual + count_google
    
    # Parámetros base para filtros (sin search_results porque ya no se usa)
    base_params = f"&last_keyword={urllib.parse.quote(last_keyword) if last_keyword else ''}&last_location={urllib.parse.quote(last_location) if last_location else ''}"
    
    html_content = f"""
    <!DOCTYPE html>
    <html lang="es">
    <head>
        <meta charset="UTF-8">
        <title>ORASIC Sales Machine</title>
        <style>
            body {{ background:#0F172A; color:white; font-family:sans-serif; padding:40px; margin:0; }}
            table {{ width:100%; border-collapse:collapse; background:#1E293B; border-radius:8px; overflow:hidden; }}
            th {{ background:#334155; color:#CBD5E1; padding:12px; text-align:left; font-size:0.85rem; }}
            tr:hover {{ background:#334155; }}
            .btn-export {{ background:#F59E0B; color:white; padding:10px 20px; text-decoration:none; border-radius:6px; font-weight:bold; display:inline-block; }}
            .search-box {{ background:#1E293B; padding:20px; border-radius:8px; margin-bottom:30px; display:flex; gap:10px; flex-wrap:wrap; align-items:end; }}
            select, input {{ padding:10px; border-radius:4px; border:1px solid #475569; background:#0F172A; color:white; flex:1; min-width:150px; }}
            button.search-btn {{ background:#A78BFA; color:white; border:none; padding:10px 20px; border-radius:4px; cursor:pointer; font-weight:bold; }}
            .other-input {{ display:none; margin-top:5px; }}
            .filter-buttons {{ display:flex; gap:10px; margin-bottom:20px; }}
            .filter-btn {{ background:#1E293B; color:#CBD5E1; padding:8px 16px; border-radius:6px; text-decoration:none; border:1px solid #475569; }}
            .filter-btn.active {{ background:#A78BFA; color:white; border-color:#A78BFA; }}
        </style>
        <script>
            function toggleOther(sId, iId) {{ 
                var s=document.getElementById(sId); 
                var i=document.getElementById(iId); 
                if(s.value==="Otro"){{i.style.display="block";i.required=true;}}
                else{{i.style.display="none";i.required=false;}} 
            }}
        </script>
    </head>
    <body>
        <div style="max-width:1400px; margin:0 auto;">
            <h1 style="text-align:center;">ORASIC Sales Machine</h1>
            {alert_html}
            
            <div class="search-box">
                <form action="/api/auto-search" method="post" style="display:flex; gap:10px; flex:1; flex-wrap:wrap;">
                    <div style="flex:1; min-width:200px;">
                        <label style="font-size:0.8rem; color:#94A3B8;">Categoría (Google Maps)</label>
                        <select name="keyword" id="keyword_select" onchange="toggleOther('keyword_select', 'keyword_other')" required>
                            <option value="" disabled selected>Selecciona...</option>
                            {rubro_options}
                        </select>
                        <input type="text" id="keyword_other" name="keyword_other" placeholder="Otro..." class="other-input">
                    </div>
                    <div style="flex:1; min-width:200px;">
                        <label style="font-size:0.8rem; color:#94A3B8;">Distrito</label>
                        <select name="location" id="location_select" onchange="toggleOther('location_select', 'location_other')" required>
                            <option value="" disabled selected>Selecciona...</option>
                            {distrito_options}
                        </select>
                        <input type="text" id="location_other" name="location_other" placeholder="Otro..." class="other-input">
                    </div>
                    <div style="width:100px;">
                        <label style="font-size:0.8rem; color:#94A3B8;">Cantidad</label>
                        <input type="number" name="limit" value="100" min="1" max="100">
                    </div>
                    <button type="submit" class="search-btn"> Buscar</button>
                </form>
            </div>
            
            <div class="filter-buttons">
                <a href="/?filter_origin=all{base_params}" class="filter-btn {'active' if filter_origin == 'all' else ''}"> Todos ({count_total})</a>
                <a href="/?filter_origin=google{base_params}" class="filter-btn {'active' if filter_origin == 'google' else ''}">[GOOGLE] ({count_google})</a>
                <a href="/?filter_origin=manual{base_params}" class="filter-btn {'active' if filter_origin == 'manual' else ''}">[MAESTRO] ({count_manual})</a>
            </div>

            <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:20px; background:#1E293B; padding:20px; border-radius:8px;">
                <div><strong>Enviados:</strong> <span style="color:#22C55E; font-size:1.5rem;">{leads_enviados}</span></div>
                <div><strong>Vista Actual:</strong> <span style="color:#F59E0B; font-size:1.5rem;">{view_title}: {len(filtered_leads)}</span></div>
                <a href="/api/export-excel" class="btn-export"> Exportar Excel</a>
            </div>
            
            <table>
                <thead><tr>
                    <th>#</th><th>Nombre</th><th>Rubro</th><th>Distrito</th><th>Contacto</th>
                    <th>Rating</th><th>Web/Redes</th><th>Plan</th><th>Origen</th><th>Acción</th>
                </tr></thead>
                <tbody>{rows_html}</tbody>
            </table>
            
            <div style="margin-top:20px; text-align:center;">
                 <a href="/" style="color:#94A3B8; text-decoration:underline;">Limpiar búsqueda</a>
            </div>

            <div style="margin-top:40px; text-align:center;">
                <p style="color:#64748B;">v12.0 • Persistencia Total en BD • Contadores Reales • Visualización Limpia</p>
            </div>
        </div>
    </body>
    </html>
    """
    return HTMLResponse(content=html_content)

@app.post("/api/mark-sent/{lead_id}")
async def mark_sent(lead_id: str):
    if not supabase_connected: return JSONResponse({"status": "error"}, status_code=500)
    try:
        headers = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}", "Content-Type": "application/json"}
        url = f"{SUPABASE_URL}/rest/v1/leads_maestros?id=eq.{lead_id}"
        resp = requests.patch(url, headers=headers, json={"Estado": "Enviado"}, timeout=5)
        return JSONResponse({"status": "success"})
    except Exception as e:
        logger.error(f"Error marcando enviado: {e}")
        return JSONResponse({"status": "error"}, status_code=500)

@app.get("/api/export-excel")
async def export_excel():
    if not supabase_connected: return JSONResponse({"error": "No conectado"}, status_code=500)
    leads = get_all_leads()
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads Maestros"
    
    headers = ["Categoria", "Distrito", "Negocio", "Direccion", "Rating", "Reseñas", "Contacto", "Web/Redes", "Publico objetivo", "Puntos fuertes", "Puntos debiles", "Presencia digital", "Responde reseñas", "Plan_sugerido"]
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        
    for lead in leads:
        ws.append([
            lead.get("Categoria", ""), lead.get("Distrito", ""), lead.get("Negocio", ""),
            lead.get("Direccion", ""), lead.get("Rating", ""), lead.get("Reseñas", ""),
            lead.get("Contacto", ""), lead.get("Web/Redes", ""), lead.get("Publico objetivo", ""),
            lead.get("Puntos fuertes", ""), lead.get("Puntos debiles", ""),
            lead.get("Presencia digital", ""), lead.get("Responde reseñas", ""),
            lead.get("Plan_sugerido", "")
        ])
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=leads_maestros_orasic.xlsx"})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)